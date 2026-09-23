import os
import json
import logging
import tempfile
import re
import time
from html import escape

from dotenv import load_dotenv

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, InputFile
from telegram.ext import (
    Application,
    CommandHandler,
    ConversationHandler,
    ContextTypes,
    MessageHandler,
    CallbackQueryHandler,
    filters,
)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from google import genai
from google.genai import types


# ============================================================
# SOZLAMALAR
# ============================================================

load_dotenv()

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

ai_client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
    handlers=[
        logging.FileHandler("bot.log", encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
logger = logging.getLogger(__name__)


# ============================================================
# HOLATLAR
# ============================================================

TOPIC, LANGUAGE, SLIDE_COUNT = range(3)

SLIDE_COUNT_OPTIONS = [8, 10, 12, 15, 20, 25]

LANGUAGES = {
    "uz": "🇺🇿 O‘zbek tili",
    "ru": "🇷🇺 Rus tili",
    "en": "🇬🇧 Ingliz tili",
}

LANGUAGE_NAMES = {
    "uz": "o‘zbek tilida",
    "ru": "rus tilida",
    "en": "ingliz tilida",
}

# Urinib ko‘riladigan modellar
MODELS_TO_TRY = [
    "gemini-3.6-flash",
    "gemini-3.5-flash",
    "gemini-3.5-flash-lite",
    "gemini-3.8-flash",
]


# ============================================================
# YORDAMCHI
# ============================================================

def to_sentence_case(text: str) -> str:
    if not text:
        return text
    text = text.strip()
    return text[0].upper() + text[1:].lower() if len(text) > 1 else text.upper()


# ============================================================
# GEMINI
# ============================================================

def generate_outline(topic: str, slide_count: int, language: str, retries: int = 2) -> dict:
    if not ai_client:
        raise RuntimeError("GEMINI_API_KEY .env faylida yo‘q!")

    language_name = LANGUAGE_NAMES.get(language, "o‘zbek tilida")

    prompt = f"""
Siz professional akademik prezentatsiyalar tayyorlovchi sun'iy intellektsiz.

Mavzu: "{topic}"
Slaydlar soni: {slide_count}
Til: {language_name}

JAVOBNI FAQAT JSON FORMATIDA QAYTARING. Markdown yoki boshqa matn yozmang.

Format:
{{
  "title": "Taqdimot nomi",
  "subtitle": "Qisqa subtitle",
  "slides": [
    {{
      "heading": "Slayd sarlavhasi",
      "bullets": [
        "Muhim ma'lumot 1",
        "Muhim ma'lumot 2",
        "Muhim ma'lumot 3",
        "Muhim ma'lumot 4"
      ]
    }}
  ],
  "sources": ["Manba 1", "Manba 2", "Manba 3"]
}}

TALABLAR:
1. Barcha matn {language_name} bo‘lsin.
2. "slides" massivida AYNAN {slide_count} ta element bo‘lsin.
3. Har bir slaydda 3-5 ta mazmunli bullet bo‘lsin.
4. Slaydlar mantiqiy ketma-ketlikda bo‘lsin.
5. Takrorlamang, faktlarni o‘ylab topmang.
6. sources da kamida 3 ta ishonchli manba yozing.
7. Sarlavhalarni oddiy yozing (Title Case qilmang).
"""

    last_error = None

    for model_name in MODELS_TO_TRY:
        for attempt in range(1, retries + 2):
            try:
                logger.info(f"Gemini so‘rovi ({model_name}, urinish {attempt}): {topic}")

                response = ai_client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        temperature=0.7,
                    ),
                )

                raw_text = response.text.strip()
                raw_text = re.sub(r"^```json\s*", "", raw_text, flags=re.IGNORECASE)
                raw_text = re.sub(r"^```\s*", "", raw_text)
                raw_text = re.sub(r"\s*```$", "", raw_text)

                data = json.loads(raw_text.strip())
                slides = data.get("slides", [])

                if len(slides) == slide_count:
                    data["title"] = to_sentence_case(data.get("title", topic))
                    data["subtitle"] = to_sentence_case(data.get("subtitle", ""))
                    for s in data["slides"]:
                        s["heading"] = to_sentence_case(s.get("heading", ""))
                    logger.info(f"✅ {slide_count} ta slayd olindi ({model_name}).")
                    return data

                last_error = f"Slayd soni mos emas: {len(slides)} != {slide_count}"
                logger.warning(last_error)

            except Exception as e:
                logger.exception(f"Xato ({model_name}): {e}")
                last_error = str(e)
                if "503" in str(e) or "UNAVAILABLE" in str(e):
                    time.sleep(3)

    logger.error(f"Barcha urinishlar muvaffaqiyatsiz: {last_error}")

    return {
        "title": to_sentence_case(topic),
        "subtitle": "Taqdimot",
        "slides": [
            {
                "heading": f"{i + 1}-slayd",
                "bullets": [
                    "Ma'lumot olishda xatolik yuz berdi.",
                    "Iltimos, qaytadan urinib ko‘ring.",
                    "API sozlamalarini tekshiring.",
                ],
            }
            for i in range(slide_count)
        ],
        "sources": [],
    }


# ============================================================
# POWERPOINT YARATISH
# ============================================================

def build_pptx(outline: dict, output_path: str, language: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    PRIMARY = RGBColor(0x1A, 0x36, 0x5D)
    ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
    LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xFE)
    DARK = RGBColor(0x1E, 0x29, 0x3B)
    GRAY = RGBColor(0x4B, 0x55, 0x63)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    SOFT = RGBColor(0xF8, 0xFA, 0xFC)
    CARD = RGBColor(0xF1, 0xF5, 0xF9)

    blank = prs.slide_layouts[6]
    FONT = "Times New Roman"

    slides_data = outline.get("slides", [])
    total = len(slides_data) + 1 + (1 if outline.get("sources") else 0)

    # ---------- TITLE SLIDE ----------
    slide = prs.slides.add_slide(blank)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY
    bg.line.fill.background()

    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.4))
    header.fill.solid()
    header.fill.fore_color.rgb = ACCENT
    header.line.fill.background()

    uni = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.5), Inches(0.5))
    utf = uni.text_frame
    up = utf.paragraphs[0]
    up.text = "TAQDIMOT"
    up.font.name = FONT
    up.font.size = Pt(18)
    up.font.bold = True
    up.font.color.rgb = WHITE
    up.alignment = PP_ALIGN.CENTER

    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = outline.get("title", "Taqdimot")
    p.font.name = FONT
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    sub = outline.get("subtitle", "")
    if sub:
        sub_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.5), Inches(10.3), Inches(0.7))
        stf = sub_box.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.text = sub
        sp.font.name = FONT
        sp.font.size = Pt(18)
        sp.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)
        sp.alignment = PP_ALIGN.CENTER

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(5.5), Inches(2.3), Inches(0.07))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    line.line.fill.background()

    # ---------- CONTENT SLIDES ----------
    for i, item in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank)

        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = WHITE
        bg.line.fill.background()

        top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
        top.fill.solid()
        top.fill.fore_color.rgb = ACCENT
        top.line.fill.background()

        left = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.14), prs.slide_height)
        left.fill.solid()
        left.fill.fore_color.rgb = PRIMARY
        left.line.fill.background()

        heading = item.get("heading", f"{i}-slayd")
        h_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12.0), Inches(0.7))
        htf = h_box.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = heading
        hp.font.name = FONT
        hp.font.size = Pt(24)
        hp.font.bold = True
        hp.font.color.rgb = DARK

        under = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.95), Inches(1.5), Inches(0.06))
        under.fill.solid()
        under.fill.fore_color.rgb = ACCENT
        under.line.fill.background()

        bullets = item.get("bullets", [])[:5]
        count = len(bullets)

        if count <= 3:
            text_width = Inches(7.2)
            start_y = 1.35

            for j, bullet in enumerate(bullets):
                y = start_y + j * 1.35

                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(0.55), Inches(y),
                    text_width, Inches(1.15)
                )
                card.fill.solid()
                card.fill.fore_color.rgb = CARD
                card.line.fill.background()

                dot = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL,
                    Inches(0.75), Inches(y + 0.4),
                    Inches(0.28), Inches(0.28)
                )
                dot.fill.solid()
                dot.fill.fore_color.rgb = ACCENT if j % 2 == 0 else RGBColor(0x0E, 0xA5, 0xE9)
                dot.line.fill.background()

                t_box = slide.shapes.add_textbox(
                    Inches(1.2), Inches(y + 0.25),
                    Inches(6.3), Inches(0.75)
                )
                ttf = t_box.text_frame
                ttf.word_wrap = True
                tp = ttf.paragraphs[0]
                tp.text = bullet
                tp.font.name = FONT
                tp.font.size = Pt(16)
                tp.font.color.rgb = DARK

            right_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(8.2), Inches(1.5),
                Inches(4.5), Inches(4.6)
            )
            right_box.fill.solid()
            right_box.fill.fore_color.rgb = LIGHT_BLUE
            right_box.line.fill.background()

            big = slide.shapes.add_textbox(Inches(8.2), Inches(2.8), Inches(4.5), Inches(1.5))
            btf = big.text_frame
            bp = btf.paragraphs[0]
            bp.text = str(i)
            bp.font.name = FONT
            bp.font.size = Pt(72)
            bp.font.bold = True
            bp.font.color.rgb = ACCENT
            bp.alignment = PP_ALIGN.CENTER

            note = slide.shapes.add_textbox(Inches(8.4), Inches(4.6), Inches(4.1), Inches(0.8))
            ntf = note.text_frame
            ntf.word_wrap = True
            np = ntf.paragraphs[0]
            np.text = heading[:40] + ("..." if len(heading) > 40 else "")
            np.font.name = FONT
            np.font.size = Pt(13)
            np.font.color.rgb = GRAY
            np.alignment = PP_ALIGN.CENTER

        else:
            if count == 4:
                card_h, gap, start_y, fsize = 0.95, 0.16, 1.25, 15
            else:
                card_h, gap, start_y, fsize = 0.82, 0.12, 1.2, 14

            for j, bullet in enumerate(bullets):
                y = start_y + j * (card_h + gap)

                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(0.55), Inches(y),
                    Inches(12.2), Inches(card_h)
                )
                card.fill.solid()
                card.fill.fore_color.rgb = CARD
                card.line.fill.background()

                dot = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL,
                    Inches(0.75), Inches(y + card_h / 2 - 0.13),
                    Inches(0.26), Inches(0.26)
                )
                dot.fill.solid()
                dot.fill.fore_color.rgb = ACCENT if j % 2 == 0 else RGBColor(0x0E, 0xA5, 0xE9)
                dot.line.fill.background()

                t_box = slide.shapes.add_textbox(
                    Inches(1.2), Inches(y + 0.18),
                    Inches(11.2), Inches(card_h - 0.3)
                )
                ttf = t_box.text_frame
                ttf.word_wrap = True
                tp = ttf.paragraphs[0]
                tp.text = bullet
                tp.font.name = FONT
                tp.font.size = Pt(fsize)
                tp.font.color.rgb = DARK

        bottom = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.15), prs.slide_width, Inches(0.35))
        bottom.fill.solid()
        bottom.fill.fore_color.rgb = SOFT
        bottom.line.fill.background()

        num_box = slide.shapes.add_textbox(Inches(11.6), Inches(7.18), Inches(1.4), Inches(0.28))
        ntf = num_box.text_frame
        npar = ntf.paragraphs[0]
        npar.text = f"{i + 1} / {total}"
        npar.font.name = FONT
        npar.font.size = Pt(11)
        npar.font.color.rgb = GRAY
        npar.alignment = PP_ALIGN.RIGHT

    # ---------- SOURCES SLIDE ----------
    sources = outline.get("sources", [])
    if sources:
        slide = prs.slides.add_slide(blank)

        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = WHITE
        bg.line.fill.background()

        top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
        top.fill.solid()
        top.fill.fore_color.rgb = ACCENT
        top.line.fill.background()

        left = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.14), prs.slide_height)
        left.fill.solid()
        left.fill.fore_color.rgb = PRIMARY
        left.line.fill.background()

        source_title = {
            "uz": "Foydalanilgan manbalar",
            "ru": "Использованные источники",
            "en": "References",
        }.get(language, "Foydalanilgan manbalar")

        h_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.65))
        htf = h_box.text_frame
        hp = htf.paragraphs[0]
        hp.text = source_title
        hp.font.name = FONT
        hp.font.size = Pt(24)
        hp.font.bold = True
        hp.font.color.rgb = DARK

        under = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.95), Inches(1.5), Inches(0.06))
        under.fill.solid()
        under.fill.fore_color.rgb = ACCENT
        under.line.fill.background()

        for idx, source in enumerate(sources[:6]):
            y = 1.4 + idx * 0.8

            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), Inches(y), Inches(0.4), Inches(0.4))
            circle.fill.solid()
            circle.fill.fore_color.rgb = ACCENT
            circle.line.fill.background()

            num = slide.shapes.add_textbox(Inches(0.7), Inches(y + 0.05), Inches(0.4), Inches(0.3))
            ntf = num.text_frame
            np = ntf.paragraphs[0]
            np.text = str(idx + 1)
            np.font.name = FONT
            np.font.size = Pt(13)
            np.font.bold = True
            np.font.color.rgb = WHITE
            np.alignment = PP_ALIGN.CENTER

            s_box = slide.shapes.add_textbox(Inches(1.35), Inches(y + 0.05), Inches(11.2), Inches(0.45))
            stf = s_box.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.text = source
            sp.font.name = FONT
            sp.font.size = Pt(15)
            sp.font.color.rgb = DARK

    prs.save(output_path)


# ============================================================
# BOT FUNKSIYALARI
# ============================================================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()
    await update.message.reply_text(
        "👋 Salom!\n\n"
        "Men AI yordamida PowerPoint taqdimotlar tayyorlab beruvchi botman.\n\n"
        "✍️ Taqdimot mavzusini yozib yuboring.\n\n"
        "Masalan:\n"
        "• Sun'iy intellekt\n"
        "• Kiberxavfsizlik\n"
        "• Bulutli texnologiyalar\n"
        "• O‘zbekiston iqtisodiyoti"
    )
    return TOPIC


async def receive_topic(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    topic = update.message.text.strip()

    if len(topic) < 3:
        await update.message.reply_text("⚠️ Mavzu juda qisqa. Aniqroq yozing.")
        return TOPIC

    context.user_data["topic"] = topic

    keyboard = [
        [InlineKeyboardButton(LANGUAGES["uz"], callback_data="lang_uz")],
        [InlineKeyboardButton(LANGUAGES["ru"], callback_data="lang_ru")],
        [InlineKeyboardButton(LANGUAGES["en"], callback_data="lang_en")],
    ]

    await update.message.reply_text(
        f"📌 Mavzu:\n<b>{escape(topic)}</b>\n\n🌐 Taqdimot tilini tanlang:",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="HTML",
    )
    return LANGUAGE


async def receive_language(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    await query.answer()

    language = query.data.replace("lang_", "")
    context.user_data["language"] = language

    keyboard = []
    row = []
    for number in SLIDE_COUNT_OPTIONS:
        row.append(InlineKeyboardButton(str(number), callback_data=f"slides_{number}"))
        if len(row) == 3:
            keyboard.append(row)
            row = []
    if row:
        keyboard.append(row)

    await query.edit_message_text(
        f"📌 Mavzu: <b>{escape(context.user_data['topic'])}</b>\n"
        f"🌐 Til: {LANGUAGES[language]}\n\n"
        "📊 Nechta slayd kerak?",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="HTML",
    )
    return SLIDE_COUNT


async def receive_slide_count(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    query = update.callback_query
    await query.answer()

    slide_count = int(query.data.replace("slides_", ""))
    topic = context.user_data["topic"]
    language = context.user_data["language"]

    await query.edit_message_text(
        f"⏳ Taqdimot tayyorlanmoqda...\n\n"
        f"📌 Mavzu: <b>{escape(topic)}</b>\n"
        f"🌐 Til: {LANGUAGES[language]}\n"
        f"📊 Slaydlar: {slide_count}\n\n"
        "Iltimos, 30–90 soniya kuting.",
        parse_mode="HTML",
    )

    try:
        outline = generate_outline(topic, slide_count, language)

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            output_path = tmp.name

        build_pptx(outline, output_path, language)

        filename = f"{topic[:40].strip().replace(' ', '_')}.pptx"

        with open(output_path, "rb") as f:
            await query.message.reply_document(
                document=InputFile(f, filename=filename),
                caption=(
                    f"✅ Tayyor!\n\n"
                    f"📌 {escape(outline.get('title', topic))}\n"
                    f"📊 {len(outline.get('slides', []))} ta slayd"
                ),
                parse_mode="HTML",
            )

        os.unlink(output_path)

    except Exception as e:
        logger.exception("Xatolik")
        await query.message.reply_text(
            f"❌ Xatolik yuz berdi:\n<code>{escape(str(e))}</code>\n\n"
            "Qaytadan /start bosing.",
            parse_mode="HTML",
        )

    context.user_data.clear()
    return ConversationHandler.END


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()
    await update.message.reply_text("❌ Bekor qilindi. /start bosing.")
    return ConversationHandler.END


# ============================================================
# ISHGA TUSHIRISH
# ============================================================

def main() -> None:
    if not TELEGRAM_BOT_TOKEN:
        raise RuntimeError("TELEGRAM_BOT_TOKEN .env faylida yo‘q!")

    app = Application.builder().token(TELEGRAM_BOT_TOKEN).build()

    conv = ConversationHandler(
        entry_points=[CommandHandler("start", start)],
        states={
            TOPIC: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_topic)],
            LANGUAGE: [CallbackQueryHandler(receive_language, pattern=r"^lang_")],
            SLIDE_COUNT: [CallbackQueryHandler(receive_slide_count, pattern=r"^slides_")],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
    )

    app.add_handler(conv)
    app.add_handler(CommandHandler("cancel", cancel))

    logger.info("Bot ishga tushdi...")
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()